"""Load schedule workbook (plan 1.10): the API's load schedule rows and totals, stamped.

The rows are the API's own load schedule (one per building or load point, columns as the API names them) and the
totals are the calc service's diversified group demand that the API shows with them. Nothing is summed here: an
ADMD column added up would not be the group's demand.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import DocumentRequest
from .common import stamp


def render(req: DocumentRequest) -> bytes:
    rows = req.rows or []
    s = stamp(req.meta, req.design)
    wb = Workbook()
    ws = wb.active
    ws.title = "Load schedule"
    ws.append(["Load schedule"])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in s.lines():
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    if not s.fit:
        ws.cell(ws.max_row, 2).font = Font(bold=True, color="C8102E")
    ws.append([])
    columns: list[str] = []
    for r in rows:
        columns += [k for k in r if k not in columns]
    if not columns:
        ws.append(["No loads in the schedule."])
    else:
        ws.append(columns)
        head = ws.max_row
        for cell in ws[head]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F2F4F6")
        for r in rows:
            ws.append([r.get(c) for c in columns])
        last = ws.max_row
        if req.totals:
            ws.append([])
            ws.append(["Totals (calc service, diversified)"])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for k, v in req.totals.items():
                ws.append([k, v])
        ws.freeze_panes = ws.cell(head + 1, 1)
        ws.auto_filter.ref = f"A{head}:{get_column_letter(len(columns))}{last}"
        for i, c in enumerate(columns, start=1):
            width = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)
    wb.properties.title = f"{s.project} load schedule {s.revision}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

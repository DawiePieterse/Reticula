"""Bill of quantities (plan 6.4): a workbook with a sheet per network part and a summary, and the same as a PDF.

Quantities, rates and amounts are the cost model's (cost.price); this only lays them out and totals them. When the
rates are the indicative library, every sheet and page says ESTIMATE.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from . import DocumentRequest
from .common import esc, money, num, pdf_text, stamp
from .report import _numbered, _status, kv, styles, table

HEAD = ["Code", "Description", "Unit", "Quantity", "Rate", "Amount", "Low", "High"]
MONEY = '#,##0.00'


def _groups(req: DocumentRequest) -> dict[str, list]:
    out: dict[str, list] = {}
    for line in req.design.cost.lines:
        out.setdefault(line.group, []).append(line)
    return out


def _title(req: DocumentRequest) -> str:
    return ("ESTIMATE: " if req.design.cost.indicative else "") + "Bill of quantities"


def render_xlsx(req: DocumentRequest) -> bytes:
    s = stamp(req.meta, req.design)
    c = req.design.cost
    wb = Workbook()
    bold, red = Font(bold=True), Font(bold=True, color="C8102E")
    band = PatternFill("solid", fgColor="F2F4F6")

    def header(ws, title: str) -> int:
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=14, color="C8102E" if c.indicative else "000000")
        for k, v in s.lines():
            ws.append([k, v])
            ws.cell(ws.max_row, 1).font = bold
        if not s.fit:
            ws.cell(ws.max_row, 2).font = red
        if c.indicative:
            ws.append(["Rates", f"ESTIMATE: indicative library '{c.library}' ({c.rate_date}), not a price"])
            ws.cell(ws.max_row, 2).font = red
        ws.append([])
        return ws.max_row + 1

    ws = wb.active
    ws.title = "Summary"
    header(ws, _title(req) + ": summary")
    ws.append(["Group", "Lines", "Amount", "Low", "High"])
    top = ws.max_row
    for cell in ws[top]:
        cell.font, cell.fill = bold, band
    for g, lines in _groups(req).items():
        ws.append([g, len(lines), sum(x.amount or 0 for x in lines), sum(x.low or 0 for x in lines), sum(x.high or 0 for x in lines)])
    ws.append(["Total", len(c.lines), c.capex, c.capex_low, c.capex_high])
    for cell in ws[ws.max_row]:
        cell.font = bold
    ws.append([])
    ws.append(["Present value of losses", None, c.losses.npv])
    ws.append(["Lifetime cost", None, c.lifetime, c.lifetime_low, c.lifetime_high])
    unpriced = [x.code for x in c.lines if x.amount is None]
    if unpriced:
        ws.append(["Not priced (no rate)", ", ".join(unpriced)])
        ws.cell(ws.max_row, 2).font = red
    for row in ws.iter_rows(min_row=top + 1, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = MONEY
    _widths(ws, [28, 70, 18, 18, 18])

    for g, lines in _groups(req).items():
        sh = wb.create_sheet(g[:31])
        header(sh, f"{_title(req)}: {g}")
        sh.append(HEAD)
        first = sh.max_row
        for cell in sh[first]:
            cell.font, cell.fill = bold, band
        for x in lines:
            sh.append([x.code, x.description, x.unit, x.qty, x.rate, x.amount, x.low, x.high])
        sh.append(["Total", None, None, None, None, sum(x.amount or 0 for x in lines), sum(x.low or 0 for x in lines), sum(x.high or 0 for x in lines)])
        for cell in sh[sh.max_row]:
            cell.font = bold
        for row in sh.iter_rows(min_row=first + 1, min_col=5, max_col=8):
            for cell in row:
                cell.number_format = MONEY
        sh.freeze_panes = sh.cell(first + 1, 1)
        sh.auto_filter.ref = f"A{first}:H{sh.max_row - 1}"
        for row in sh.iter_rows(min_row=first + 1, min_col=2, max_col=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        _widths(sh, [26, 60, 10, 12, 14, 16, 16, 16])
    wb.properties.title = f"{s.project} {_title(req)} {s.revision}"
    wb.properties.subject = s.status
    wb.properties.creator = "Reticula calc service"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def render_pdf(req: DocumentRequest) -> bytes:
    d = req.design
    s = stamp(req.meta, d)
    c = d.cost
    st = styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=17 * mm,
                            title=pdf_text(f"{s.project} {_title(req)} {s.revision}"), subject=pdf_text(s.status), creator="Reticula calc service")
    story: list = [Paragraph(esc(_title(req)), st["h1"]), Paragraph(esc(s.project), st["h2"]), kv(s.lines()[1:-1], st, 260 * mm),
                   Spacer(1, 3 * mm), _status(s, st)]
    if c.indicative:
        story.append(Paragraph(esc(f"ESTIMATE: indicative rates from '{c.library}' ({c.rate_date}). Not a price."), st["h3"]))
    w = [34 * mm, 104 * mm, 14 * mm, 20 * mm, 22 * mm, 24 * mm, 24 * mm, 24 * mm]
    for g, lines in _groups(req).items():
        story.append(Paragraph(esc(g), st["h2"]))
        rows = [HEAD] + [[x.code, x.description, x.unit, num(x.qty), money(x.rate, ""), money(x.amount, ""), money(x.low, ""), money(x.high, "")]
                         for x in lines]
        rows.append(["Total", "", "", "", "", money(sum(x.amount or 0 for x in lines), ""), money(sum(x.low or 0 for x in lines), ""),
                     money(sum(x.high or 0 for x in lines), "")])
        story.append(table(rows, st, w, fails=[i for i, x in enumerate(lines) if x.amount is None]))
    story += [Paragraph("Totals", st["h2"]), kv([("Capital cost", (f"{money(c.capex, c.currency)} ({money(c.capex_low, c.currency)} to "
                                                                     f"{money(c.capex_high, c.currency)})")),
                                                   ("Present value of losses", money(c.losses.npv, c.currency)),
                                                   ("Lifetime cost", money(c.lifetime, c.currency))], st, 260 * mm)]
    doc.build(story, canvasmaker=_numbered(s.footer(), not s.fit))
    return buf.getvalue()

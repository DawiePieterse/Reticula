"""Bill of quantities (plan 6.4): PDF and Excel, every figure marked as an indicative estimate.

Sections per LV site (the documented construction) and the MV network, from the stored cost lines; a summary adds
like items across sections; assembly components and any rate overrides (with their dates) are listed.
"""

from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from ..lv.costs import rates_of
from .package import DocumentPackage, chosen_option, stamp_lines, stamp_text
from .report import BODY, FONT, H1, H2, WARN, table

ESTIMATE = "ESTIMATE – indicative rates for comparing options; not a tender price."


def sections(pkg: DocumentPackage) -> list[tuple[str, list[dict], str]]:
    """Cost sections. With an MV design, transformers are costed as the MV design's units (pole-mount or mini-sub,
    including structure and MV fuses), so the LV sections leave out their bare transformer line."""
    out: list[tuple[str, list[dict], str]] = []
    mv = pkg.mv_design.result if pkg.mv_design else None
    for site in pkg.lv_designs:
        o = chosen_option(site)
        if o:
            lines = [x for x in o["cost"]["lines"] if not (mv and x["item"].startswith("Transformer"))]
            out.append((f"LV – {site.label} ({o['construction']})", lines, o["cost"]["currency"]))
    if mv:
        lines = [x for x in mv.get("cost_lines", []) if x["item"].startswith("MV") or re.match(r"Site .+: (pole-mount|minisub) ", x["item"])]
        if lines:
            out.append(("MV network and transformer units", lines, mv.get("currency", "ZAR")))
    return out


def summary(secs: list[tuple[str, list[dict], str]]) -> list[dict]:
    acc: dict[tuple[str, str, float], dict] = {}
    for _, lines, _ in secs:
        for x in lines:
            k = (x["item"], x["unit"], x["rate"])
            row = acc.setdefault(k, {"item": x["item"], "unit": x["unit"], "rate": x["rate"], "quantity": 0.0, "amount": 0.0})
            row["quantity"] += x["quantity"]
            row["amount"] += x["amount"]
    return sorted(acc.values(), key=lambda r: r["item"])


def overrides(pkg: DocumentPackage) -> list[dict]:
    return list(rates_of(pkg.rates).get("overrides", [])) if isinstance(pkg.rates, dict) else []


def render_boq_pdf(pkg: DocumentPackage) -> bytes:
    secs = sections(pkg)
    buf = io.BytesIO()
    footer = stamp_text(pkg)

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont(FONT, 6.5)
        canvas.drawString(15 * mm, 8 * mm, footer[:170])
        canvas.drawRightString(A4[0] - 15 * mm, 12 * mm, f"Page {doc.page}")
        canvas.drawString(15 * mm, A4[1] - 9 * mm, f"{pkg.project.name} – bill of quantities – {ESTIMATE}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=17 * mm,
                            title=f"{pkg.project.name} bill of quantities", author="Reticula", subject=footer)
    W = A4[0] - 30 * mm
    widths = [W - 95 * mm, 22 * mm, 15 * mm, 25 * mm, 33 * mm]
    s: list[Any] = [Paragraph(f"{pkg.project.name} – bill of quantities", H1), Paragraph(ESTIMATE, WARN), Spacer(1, 6)]
    s += [Paragraph(x, BODY) for x in stamp_lines(pkg)]
    total = 0.0
    currency = secs[0][2] if secs else "ZAR"
    for title, lines, cur in secs:
        sub = sum(x["amount"] for x in lines)
        total += sub
        s += [Paragraph(title, H2), table([["Item", "Quantity", "Unit", "Rate", f"Amount ({cur})"]] +
                                          [[x["item"], f"{x['quantity']:,.2f}", x["unit"], f"{x['rate']:,.2f}", f"{x['amount']:,.2f}"] for x in lines] +
                                          [["Subtotal", "", "", "", f"{sub:,.2f}"]], widths)]
    s += [Paragraph("Summary", H2), table([["Item", "Quantity", "Unit", "Rate", f"Amount ({currency})"]] +
                                          [[r["item"], f"{r['quantity']:,.2f}", r["unit"], f"{r['rate']:,.2f}", f"{r['amount']:,.2f}"] for r in summary(secs)] +
                                          [["Total (excluding VAT)", "", "", "", f"{total:,.2f}"]], widths)]
    comps = [(x["item"], c) for _, lines, _ in secs for x in lines for c in x.get("components", [])]
    if comps:
        seen: set[tuple[str, str]] = set()
        rows = [["Assembly", "Material", "Description", "Qty", "Rate"]]
        for item, c in comps:
            if (item, c["material"]) not in seen:
                seen.add((item, c["material"]))
                rows.append([item, c["material"], c["description"], f"{c['quantity']:g} {c['unit']}", f"{c['rate']:,.2f}"])
        s += [Paragraph("Assemblies", H2), table(rows, [35 * mm, 30 * mm, W - 110 * mm, 20 * mm, 25 * mm])]
    if ov := overrides(pkg):
        s += [Paragraph("Rate overrides", H2), table([["Section", "Code", "Rate", "Date", "Source"]] +
                                                     [[o.get("section"), o.get("code"), o.get("rate"), o.get("date"), o.get("source", "")] for o in ov])]
    doc.build(s, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


def render_boq_xlsx(pkg: DocumentPackage) -> bytes:
    secs = sections(pkg)
    wb = Workbook()
    ws = wb.active
    ws.title = "BoQ"
    bold, warn = Font(bold=True), PatternFill("solid", fgColor="FFF7ED")
    ws.append([f"{pkg.project.name} – bill of quantities"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([ESTIMATE])
    ws["A2"].fill = warn
    for line in stamp_lines(pkg):
        ws.append([line])
    ws.append([])
    ws.append(["Section", "Item", "Quantity", "Unit", "Rate", "Amount"])
    for c in ws[ws.max_row]:
        c.font = bold
    first = ws.max_row + 1
    for title, lines, _ in secs:
        for x in lines:
            ws.append([title, x["item"], x["quantity"], x["unit"], x["rate"], None])
            r = ws.max_row
            ws.cell(r, 6, f"=C{r}*E{r}")
    last = ws.max_row
    ws.append(["Total (excluding VAT)", None, None, None, None, f"=SUM(F{first}:F{last})"])
    for c in ws[ws.max_row]:
        c.font = bold
    for col, w in zip("ABCDEF", (36, 44, 12, 8, 12, 16), strict=True):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=first, min_col=3, max_col=6):
        for c in row:
            c.number_format = "#,##0.00"
            c.alignment = Alignment(horizontal="right")

    sm = wb.create_sheet("Summary")
    sm.append(["Item", "Quantity", "Unit", "Rate", "Amount"])
    for r in summary(secs):
        sm.append([r["item"], round(r["quantity"], 2), r["unit"], r["rate"], round(r["amount"], 2)])
    sm.column_dimensions["A"].width = 44

    asm = wb.create_sheet("Assemblies")
    asm.append(["Assembly", "Material", "Description", "Quantity", "Unit", "Rate"])
    seen: set[tuple[str, str]] = set()
    for _, lines, _ in secs:
        for x in lines:
            for c in x.get("components", []):
                if (x["item"], c["material"]) not in seen:
                    seen.add((x["item"], c["material"]))
                    asm.append([x["item"], c["material"], c["description"], c["quantity"], c["unit"], c["rate"]])
    if ov := overrides(pkg):
        o = wb.create_sheet("Rate overrides")
        o.append(["Section", "Code", "Rate", "Date", "Source"])
        for x in ov:
            o.append([x.get("section"), x.get("code"), x.get("rate"), x.get("date"), x.get("source", "")])
    st = wb.create_sheet("Stamp")
    for line in stamp_lines(pkg):
        st.append([line])
    wb.properties.title = f"{pkg.project.name} bill of quantities"
    wb.properties.subject = stamp_text(pkg)
    wb.properties.creator = "Reticula"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

"""Load schedule spreadsheet, registers, authority checklist and the submission pack (plans 1.10, 6.6, 6.8), and the
render entry point that makes every requested document from one package."""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import zipfile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .boq import render_boq_pdf, render_boq_xlsx
from .dxf import Sheet, render_drawings
from .gis import render_geojson, render_kml, render_shapefiles
from .package import (
    ChecklistItem,
    DocKind,
    DocumentPackage,
    RenderedFile,
    RenderRequest,
    RenderResult,
    chosen_option,
    load_template,
    stamp_lines,
    stamp_text,
    unverified,
)
from .report import BODY, FONT, H1, H2, render_report, table

TITLES: dict[str, str] = {
    "drawings": "Network drawings (DXF)", "report": "Design report", "boq_pdf": "Bill of quantities (PDF)", "boq_xlsx": "Bill of quantities (Excel)",
    "gis_geojson": "GIS data (GeoJSON)", "gis_kml": "GIS data (KML)", "gis_shp": "GIS data (Shapefile)", "loads_xlsx": "Load schedule (Excel)",
    "registers": "Drawing and document registers, authority checklist", "submission_pack": "Submission pack",
}


def render_loads_xlsx(pkg: DocumentPackage) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Load schedule"
    ws.append([f"{pkg.project.name} – load schedule"])
    ws["A1"].font = Font(bold=True, size=14)
    for line in stamp_lines(pkg):
        ws.append([line])
    ws.append([])
    ws.append(["Erf", "Building", "Kind", "Load class", "ADMD kVA", "Phases", "Status", "Overridden", "Override reason", "Longitude", "Latitude"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    first = ws.max_row + 1
    for r in sorted(pkg.loads, key=lambda r: (r.erf or "", r.building_id)):
        ws.append([r.erf, r.building_id, r.kind, r.load_class, r.kva, r.phases, r.status, "yes" if r.overridden else "", r.override_reason, r.lon, r.lat])
    ws.append(["Total", None, None, None, f"=SUM(E{first}:E{ws.max_row})"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for col, width in zip("ABCDEFGHIJK", (10, 38, 12, 22, 10, 8, 12, 10, 30, 12, 12), strict=True):
        ws.column_dimensions[col].width = width
    wb.properties.subject = stamp_text(pkg)
    wb.properties.creator = "Reticula"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def checklist(pkg: DocumentPackage) -> list[ChecklistItem]:
    """The authority's submission checklist, with what Reticula can tell from the package."""
    tpl = load_template(pkg.project.authority)
    lv_ok = bool(pkg.lv_designs) and all((o := chosen_option(s)) is not None and o.get("passed") for s in pkg.lv_designs)
    facts: dict[str, tuple[bool, str]] = {
        "connection_point": (bool(pkg.connection_point and pkg.connection_point.get("available_capacity_kva") and pkg.connection_point.get("fault_mva_max")),
                             f"reference {pkg.connection_point.get('reference') or '—'}" if pkg.connection_point else "no connection point entered"),
        "loads_recorded": (bool(pkg.loads) and all(r.kva is not None for r in pkg.loads), f"{len(pkg.loads)} loads in the schedule"),
        "lv_passed": (lv_ok, f"{len(pkg.lv_designs)} LV site(s)"),
        "mv_passed": (bool(pkg.mv_design and pkg.mv_design.result.get("passed")), "no MV design" if not pkg.mv_design else ""),
        "bulk_passed": (bool(pkg.bulk_study and pkg.bulk_study.result.get("passed")),
                        f"NMD {pkg.bulk_study.result.get('notified_max_demand_kva'):g} kVA" if pkg.bulk_study else "no bulk supply study"),
        "rules_verified": (not unverified(pkg), ", ".join(unverified(pkg))),
        "assumptions_cleared": (all(a.status != "open" for a in pkg.assumptions), f"{sum(1 for a in pkg.assumptions if a.status == 'open')} open"),
        "same_revision": (True, f"all generated together as revision {pkg.stamp.revision}"),
    }
    out = []
    for item in tpl.get("checklist", []):
        auto = item.get("auto", "manual")
        if auto == "manual" or auto not in facts:
            out.append(ChecklistItem(id=item["id"], text=item["text"], status="manual", detail="to be confirmed by the engineer"))
        else:
            ok, detail = facts[auto]
            out.append(ChecklistItem(id=item["id"], text=item["text"], status="met" if ok else "not met", detail=detail))
    return out


def render_registers(pkg: DocumentPackage, sheets: list[Sheet], documents: list[tuple[str, str]], items: list[ChecklistItem]) -> tuple[bytes, bytes]:
    """Registers and checklist as PDF and CSV."""
    buf = io.BytesIO()
    footer = stamp_text(pkg)

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont(FONT, 6.5)
        canvas.drawString(15 * mm, 8 * mm, footer[:170])
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
                            title=f"{pkg.project.name} registers", author="Reticula", subject=footer)
    W = A4[0] - 30 * mm
    s: list[Any] = [Paragraph(f"{pkg.project.name} – submission registers", H1)] + [Paragraph(x, BODY) for x in stamp_lines(pkg)] + [Spacer(1, 6)]
    s += [Paragraph("Drawing register", H2), table([["Number", "Title", "Scale", "Revision"]] + [[x.number, x.title, f"1:{x.scale}", pkg.stamp.revision] for x in sheets],
                                                   [45 * mm, W - 90 * mm, 22 * mm, 23 * mm])]
    s += [Paragraph("Document register", H2), table([["File", "Title", "Revision"]] + [[n, t, pkg.stamp.revision] for n, t in documents], [70 * mm, W - 93 * mm, 23 * mm])]
    s += [Paragraph("Authority checklist", H2), table([["Item", "Requirement", "Status", "Detail"]] + [[i.id, i.text, i.status, i.detail] for i in items],
                                                      [12 * mm, W - 82 * mm, 20 * mm, 50 * mm])]
    doc.build(s, onFirstPage=on_page, onLaterPages=on_page)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["register", "number_or_file", "title", "scale_or_status", "revision_or_detail"])
    for x in sheets:
        w.writerow(["drawing", x.number, x.title, f"1:{x.scale}", pkg.stamp.revision])
    for n, t in documents:
        w.writerow(["document", n, t, "", pkg.stamp.revision])
    for i in items:
        w.writerow(["checklist", i.id, i.text, i.status, i.detail])
    return buf.getvalue(), out.getvalue().encode("utf-8")


def _slug(text: str) -> str:
    return "".join(c if c.isalnum() else "-" for c in text).strip("-").lower()[:40] or "project"


def _file(kind: DocKind, name: str, content_type: str, data: bytes) -> RenderedFile:
    return RenderedFile(kind=kind, name=name, title=TITLES[kind], content_type=content_type, size=len(data), sha256=hashlib.sha256(data).hexdigest(),
                        data_b64=base64.b64encode(data).decode("ascii"))


def render(req: RenderRequest) -> RenderResult:
    pkg = req.package
    kinds = set(req.kinds)
    if "submission_pack" in kinds:
        kinds |= {"drawings", "report", "boq_pdf", "boq_xlsx", "gis_geojson", "gis_kml", "gis_shp", "loads_xlsx", "registers"}
    base = f"{_slug(pkg.project.reference or pkg.project.name)}-{_slug(pkg.stamp.revision)}"
    warnings: list[str] = []
    if unverified(pkg):
        warnings.append("The design uses rules values not yet verified against the standards; documents are marked not for submission.")
    if not pkg.lv_designs:
        warnings.append("No LV design in the package.")
    files: list[RenderedFile] = []
    raw: dict[str, tuple[str, bytes]] = {}
    sheets: list[Sheet] = []
    dxf, sheets, crs = render_drawings(pkg)
    if "drawings" in kinds:
        files.append(_file("drawings", f"{base}-drawings.dxf", "image/vnd.dxf", dxf))
        warnings.append(f"Drawings are in {crs} metres.")
    makers = [
        ("report", "report.pdf", "application/pdf", lambda: render_report(pkg, sheets)),
        ("boq_pdf", "boq.pdf", "application/pdf", lambda: render_boq_pdf(pkg)),
        ("boq_xlsx", "boq.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", lambda: render_boq_xlsx(pkg)),
        ("gis_geojson", "network.geojson", "application/geo+json", lambda: render_geojson(pkg)),
        ("gis_kml", "network.kml", "application/vnd.google-earth.kml+xml", lambda: render_kml(pkg)),
        ("gis_shp", "shapefiles.zip", "application/zip", lambda: render_shapefiles(pkg)),
        ("loads_xlsx", "load-schedule.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", lambda: render_loads_xlsx(pkg)),
    ]
    for kind, suffix, ctype, make in makers:
        if kind in kinds:
            files.append(_file(kind, f"{base}-{suffix}", ctype, make()))  # type: ignore[arg-type]
    items = checklist(pkg)
    if "registers" in kinds:
        docs = [(f.name, f.title) for f in files] + [(f"{base}-registers.pdf", TITLES["registers"])]
        pdf, csv_bytes = render_registers(pkg, sheets, docs, items)
        files.append(_file("registers", f"{base}-registers.pdf", "application/pdf", pdf))
        raw["registers_csv"] = (f"{base}-registers.csv", csv_bytes)
    if "submission_pack" in kinds:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for f in files:
                z.writestr(f.name, base64.b64decode(f.data_b64))
            for name, data in raw.values():
                z.writestr(name, data)
            z.writestr("STAMP.txt", "\n".join([pkg.project.name, *stamp_lines(pkg)]) + "\n")
        files.append(_file("submission_pack", f"{base}-submission-pack.zip", "application/zip", buf.getvalue()))
    wanted = set(req.kinds)
    return RenderResult(files=[f for f in files if f.kind in wanted], checklist=items, warnings=warnings)

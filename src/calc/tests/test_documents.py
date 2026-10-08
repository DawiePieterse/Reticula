"""Design documents (Phase 6): every renderer produces a valid file that carries the stamp and the design's own values."""

import base64
import csv
import io
import json
import zipfile
from xml.etree import ElementTree as ET

import ezdxf
import pytest
import shapefile
from ezdxf import recover
from openpyxl import load_workbook
from pypdf import PdfReader
from test_design import RULES, cp, layout

from reticula_calc.design.run import DesignRequest, run
from reticula_calc.documents import KINDS, DocumentMeta, DocumentRequest, PackFile, PackRequest, ReportSection, render
from reticula_calc.documents import pack as pack_mod
from reticula_calc.documents.common import NOT_FIT, filename, pdf_text, stamp
from reticula_calc.rules import load_rules

META = DocumentMeta(project_name="Soshanguve Ext 19", project_code="SOS-19", authority="Eskom", revision_number=2, revision_label="for comment",
                    design_date="2026-10-06", document_number="SOS-19-DWG-001", engineer_name="A Engineer", engineer_registration="20100123")


@pytest.fixture(scope="module", params=["overhead", "underground"])
def design(request):
    cands, loads, classes = layout(proposed=True)
    return run(DesignRequest(rules=RULES, candidates=cands, loads=loads, classes=classes, connection_point=cp(),
                             options={"construction": request.param}), load_rules(RULES))


def req(d, **kw) -> DocumentRequest:
    return DocumentRequest(meta=META, design=d, **kw)


def pdf_strings(data: bytes) -> str:
    """The text of every page, as a reader extracts it."""
    return " ".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)


def test_stamp_reasons_and_file_names(design):
    s = stamp(META, design)
    assert not s.fit and s.status.startswith(NOT_FIT)
    assert any("placeholder" in r for r in s.reasons) and any("not inspected" in r for r in s.reasons)
    assert s.rules == f"{RULES} ({design.rules_hash})" and "estimate only" in s.rates and s.inputs == design.inputs_hash
    assert filename(META, "report_pdf") == "sos-19_report_R2.pdf"
    assert filename(META, "pack") == "sos-19_submission-pack_R2.zip"
    signed = stamp(META.model_copy(update={"signed_off": True, "signed_off_at": "2026-10-07"}), design.model_copy(update={"fit_to_submit": True}))
    assert signed.fit and signed.status == "Fit to submit, signed off by A Engineer, ECSA 20100123 on 2026-10-07"


def test_pdf_text_spells_greek_the_font_lacks():
    assert pdf_text("ΔV = I·z·cosφ, Σ μ") == "deltaV = I·z·cosphi, Sigma mu".replace("deltaV", "DeltaV")


def test_drawing_has_the_drawing_practice_layers_symbols_and_sheet(design):
    data = render("drawing_dxf", req(design))
    doc, auditor = recover.read(io.BytesIO(data))
    assert not auditor.has_errors
    msp = doc.modelspace()
    by_layer: dict[str, int] = {}
    for e in msp:
        by_layer[e.dxf.layer] = by_layer.get(e.dxf.layer, 0) + 1
    lv = "RET-LV-UG" if design.construction == "underground" else "RET-LV-OH"
    assert by_layer[lv] == len(design.lv.network.branches)
    assert by_layer["RET-MV-OH"] == len(design.mv_network.branches)
    assert by_layer["RET-SERVICE-UG" if design.construction == "underground" else "RET-SERVICE"] == len(design.lv.allocation.allocations)
    assert by_layer["RET-ZONE"] >= 2  # the zone outline and its label
    inserts = {e.dxf.name for e in msp.query("INSERT")}
    assert {"RET_MINISUB" if design.construction == "underground" else "RET_TX", "RET_CP", "RET_POLE_MV"} <= inserts
    assert ("RET_KIOSK" if design.construction == "underground" else "RET_POLE_LV") in inserts
    assert doc.layers.get("RET-MV-OH").dxf.linetype == "RET_MV" and doc.layers.get("RET-MV-OH").dxf.color == 1
    assert doc.layers.get("RET-LV-OH").dxf.linetype == "RET_LV" and doc.layers.get("RET-LV-OH").dxf.color == 3
    assert doc.layers.get("RET-MV-UG").dxf.color == 6 and doc.layers.get("RET-LV-UG").dxf.color == 5
    texts = " ".join(e.dxf.text for e in msp.query("TEXT"))
    assert f"TX1 {design.transformers.transformers[0].rating_kva:g} kVA" in texts
    sheet = doc.layouts.get("A1 sheet")
    sheet_text = " ".join(e.dxf.text for e in sheet.query("TEXT"))
    assert NOT_FIT in sheet_text and design.rules_hash in sheet_text and "SOS-19-DWG-001" in sheet_text and "Legend" in sheet_text
    assert len(sheet.query("VIEWPORT")) >= 1
    assert "Layout1" not in [lay.name for lay in doc.layouts]


def test_report_carries_the_stamp_every_section_and_the_traceability(design):
    d = design
    data = render("report_pdf", req(d, sections=[ReportSection(title="Scope", text="The township of 22 stands.")]))
    assert data.startswith(b"%PDF")
    text = pdf_strings(data)
    for needle in (NOT_FIT, "Summary", "Engineer's notes", "The township of 22 stands.", "LV network", "Transformers", "MV network",
                   "Bulk supply", "Cost", "Appendix A. Checks", "Appendix B. Traceability", d.rules_hash, d.inputs_hash, "page 1 of"):
        assert needle in text, needle
    assert ("Underground cables" in text) == (d.construction == "underground")
    assert ("Overhead line, LV" in text) == (d.construction == "overhead")
    # Every traced value is listed with its formula id.
    for fid in {t.formula_id for t in [d.transformers.transformers[0].trace, d.cost.capex_trace, d.bulk.fault_trace] if t}:
        assert fid in text


def test_boq_workbook_totals_the_priced_lines_and_says_estimate(design):
    wb = load_workbook(io.BytesIO(render("boq_xlsx", req(design))))
    groups = {line.group for line in design.cost.lines}
    assert set(wb.sheetnames) == {"Summary", *groups}
    summary = wb["Summary"]
    assert str(summary["A1"].value).startswith("ESTIMATE")
    rows = {r[0]: r for r in summary.iter_rows(values_only=True) if r and r[0]}
    assert rows["Total"][2] == pytest.approx(design.cost.capex, abs=0.01)
    assert rows["Total"][3] == pytest.approx(design.cost.capex_low, abs=0.01)
    for g in groups:
        sheet = [r for r in wb[g].iter_rows(values_only=True) if r and r[0]]
        lines = [r for r in sheet if r[0].startswith("A-")]
        assert len(lines) == sum(1 for x in design.cost.lines if x.group == g)
        assert next(r for r in sheet if r[0] == "Total")[5] == pytest.approx(sum(x.amount or 0 for x in design.cost.lines if x.group == g), abs=0.01)


def test_boq_pdf(design):
    text = pdf_strings(render("boq_pdf", req(design)))
    assert "ESTIMATE" in text and design.cost.lines[0].code in text and NOT_FIT in text


def test_load_schedule_keeps_the_rows_and_the_calc_totals(design):
    rows = [{"erf": "101", "kind": "residential", "admd_kva": 2.37}, {"erf": "102", "kind": "special", "admd_kva": 25.0, "note": "school"}]
    wb = load_workbook(io.BytesIO(render("load_schedule_xlsx", req(design, rows=rows, totals={"Total kVA (diversified)": 26.1}))))
    ws = wb.active
    values = [r for r in ws.iter_rows(values_only=True)]
    head = next(i for i, r in enumerate(values) if r[0] == "erf")
    assert values[head][:4] == ("erf", "kind", "admd_kva", "note")
    assert values[head + 2][3] == "school"
    assert ("Total kVA (diversified)", 26.1) in [r[:2] for r in values]
    assert ws.freeze_panes == f"A{head + 2}"


def test_geojson_kml_and_shapefile_hold_the_same_features(design):
    fc = json.loads(render("geojson", req(design)))
    kinds = {f["properties"]["kind"] for f in fc["features"]}
    assert {"transformer" if design.construction == "overhead" else "minisub", "service", "connection_point", "mv_line"} <= kinds
    assert fc["metadata"]["rules"] == f"{RULES} ({design.rules_hash})" and fc["metadata"]["fit_to_submit"] is False
    n = len(fc["features"])
    root = ET.fromstring(render("kml", req(design)))
    ns = {"k": "http://www.opengis.net/kml/2.2"}
    assert len(root.findall(".//k:Placemark", ns)) == n
    z = zipfile.ZipFile(io.BytesIO(render("shapefile_zip", req(design))))
    total = 0
    for layer in ("points", "lines"):
        r = shapefile.Reader(shp=io.BytesIO(z.read(f"sos-19_{layer}.shp")), shx=io.BytesIO(z.read(f"sos-19_{layer}.shx")),
                             dbf=io.BytesIO(z.read(f"sos-19_{layer}.dbf")))
        total += len(r)
        assert "GCS_WGS_1984" in z.read(f"sos-19_{layer}.prj").decode()
    assert total == n


def test_pack_registers_every_file_and_the_checklist_reads_the_checks(design):
    files = [PackFile(name=filename(META, k), title=KINDS[k][0], kind=k, number=f"SOS-19-{i:03d}",
                      content_base64=base64.b64encode(render(k, req(design, rows=[]))).decode()) for i, k in enumerate(("drawing_dxf", "report_pdf", "boq_xlsx"), 1)]
    z = zipfile.ZipFile(io.BytesIO(pack_mod.render(PackRequest(meta=META, design=design, files=files))))
    names = z.namelist()
    assert all(f"documents/{f.name}" in names for f in files)
    dreg = list(csv.reader(io.StringIO(z.read("registers/sos-19_drawing-register_R2.csv").decode())))
    oreg = list(csv.reader(io.StringIO(z.read("registers/sos-19_document-register_R2.csv").decode())))
    assert [r[0] for r in dreg[1:]] == ["SOS-19-001"] and sorted(r[0] for r in oreg[1:]) == ["SOS-19-002", "SOS-19-003"]
    chk = {r[0]: r[1] for r in csv.reader(io.StringIO(z.read("registers/sos-19_authority-checklist_R2.csv").decode()))}
    assert chk["Every element inspected in the field"] == "fail (1)"
    assert chk["Rules values confirmed (no placeholders)"].startswith("fail")
    assert chk["Signed off by the registered engineer"] == "fail"
    assert chk["LV voltage drop within the limit at every node and service"].startswith("pass")
    # The MV line is overhead in both constructions, so its poles are always checked.
    assert chk["Pole class and stays"] == f"pass ({sum(1 for c in design.checks if c.category == 'oh_pole')})"
    assert z.read("registers/sos-19_authority-checklist_R2.pdf").startswith(b"%PDF")


def test_checklist_without_the_connection_point_says_the_bulk_studies_did_not_run():
    cands, loads, classes = layout()
    d = run(DesignRequest(rules=RULES, candidates=cands, loads=loads, classes=classes, options={"construction": "overhead"}), load_rules(RULES))
    items = {i: (s, n) for i, s, n in pack_mod.checklist(d, META)}
    assert items["Fault level within the switchgear rating"][0] == "not run"
    assert items["Connection point data from the authority (capacity, fault level)"][0] == "fail"


def test_endpoints(client, design):
    body = req(design).model_dump(mode="json")
    r = client.post("/calc/documents/report_pdf", json=body)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"] == 'attachment; filename="sos-19_report_R2.pdf"'
    assert client.post("/calc/documents/teapot", json=body).status_code == 404
    assert client.post("/calc/documents/load_schedule_xlsx", json=body).status_code == 422
    files = [{"name": "a.pdf", "title": "Design report", "kind": "report_pdf", "content_base64": base64.b64encode(b"%PDF-").decode()}]
    p = client.post("/calc/documents/pack", json={"meta": body["meta"], "design": body["design"], "files": files})
    assert p.status_code == 200 and p.headers["content-type"] == "application/zip"
    assert "documents/a.pdf" in zipfile.ZipFile(io.BytesIO(p.content)).namelist()
    assert ezdxf  # imported for recover


def test_services_and_fuses_reach_every_document():
    cands, loads, classes = layout()
    d = run(DesignRequest(rules="eskom/0.9.0", candidates=cands, loads=loads, classes=classes, connection_point=cp(),
                          options={"construction": "overhead"}), load_rules("eskom/0.9.0"))
    # One service pole halfway along the first service, as the design places it where the sag is too low.
    sv = d.services.services[0]
    a = next(x for x in d.lv.allocation.allocations if x.load_id == sv.load_id)
    sv.poles = [((a.at[0] + a.location[0]) / 2, (a.at[1] + a.location[1]) / 2)]
    d.services.service_poles = 1
    doc, _ = recover.read(io.BytesIO(render("drawing_dxf", req(d))))
    msp = doc.modelspace()
    assert len([e for e in msp.query("INSERT") if e.dxf.layer == "RET-SERVICE"]) == 1
    assert "SP1 7m" in " ".join(e.dxf.text for e in msp.query("TEXT"))
    fc = json.loads(render("geojson", req(d)))
    assert [f["properties"]["label"] for f in fc["features"] if f["properties"]["kind"] == "service_pole"] == ["SP1"]
    assert {f["properties"]["conductor"] for f in fc["features"] if f["properties"]["kind"] == "service"} == {"AIRDAC-SNE-10"}
    text = pdf_strings(render("report_pdf", req(d)))
    for needle in ("Services", "AIRDAC-SNE-10", "Fuse A", "Service poles added", "lv.protection.fuse.v1", "oh.service.span.v1"):
        assert needle in text, needle
    items = {i: s for i, s, _ in pack_mod.checklist(d, META)}
    assert items["LV feeder fuse between the design current and the conductor rating"].startswith("pass")
    assert items["Service cable drop within the service limit"].startswith("pass")
    assert items["Overhead service clearance, with service poles where needed"].startswith("pass")
    assert items["LV fault current at feeder ends"].startswith("pass")

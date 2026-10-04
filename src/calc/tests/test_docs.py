"""Design documents (plan Phase 6): every document is made from the package, carries the stamp and reads back."""

import base64
import io
import json
import zipfile
from xml.dom import minidom

import ezdxf
import openpyxl
import pytest
import shapefile
from docs_fixture import package
from pypdf import PdfReader

from reticula_calc.docs.boq import sections
from reticula_calc.docs.pack import render
from reticula_calc.docs.package import RenderRequest


@pytest.fixture(scope="module")
def rendered():
    return render(RenderRequest(package=package()))


def data(rendered, kind: str) -> bytes:
    return base64.b64decode(next(f for f in rendered.files if f.kind == kind).data_b64)


def pdf_text(raw: bytes) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(raw)).pages)


def test_every_kind_is_made_with_a_hash(rendered):
    kinds = [f.kind for f in rendered.files]
    assert kinds == ["drawings", "report", "boq_pdf", "boq_xlsx", "gis_geojson", "gis_kml", "gis_shp", "loads_xlsx", "registers", "submission_pack"]
    assert all(f.size == len(base64.b64decode(f.data_b64)) and len(f.sha256) == 64 for f in rendered.files)
    assert all(f.name.startswith("tt1-d1-") for f in rendered.files)


def test_drawings_have_layers_sheets_and_the_stamp(rendered):
    doc = ezdxf.read(io.StringIO(data(rendered, "drawings").decode("utf-8")))
    assert not doc.audit().has_errors
    layers = {e.dxf.layer for e in doc.modelspace()}
    assert {"LV-FEEDER-OH", "LV-SERVICE", "LV-POLE", "TRANSFORMER", "MV-LINE-OH", "CONNECTION-POINT", "STANDS"} <= layers
    sheets = [lay for lay in doc.layouts if lay.name != "Model"]
    assert [s.name for s in sheets] == ["RET-TT1-001", "RET-TT1-002", "RET-TT1-003"]
    texts = [t.dxf.text for t in sheets[1].query("TEXT")]
    assert "RET-TT1-002" in texts and "D1" in texts and any("eskom/0.5.0" in t for t in texts)
    assert any("NOT FOR SUBMISSION" in t for t in texts)
    vp = next(v for v in sheets[1].query("VIEWPORT") if v.dxf.id > 1)
    assert vp.dxf.view_height == pytest.approx(574 * 500 / 1000)  # 1:500 on A1
    pts = [p for e in doc.modelspace().query("LWPOLYLINE[layer=='LV-FEEDER-OH']") for p in e.get_points("xy")]
    assert all(-350_000 < x < 350_000 and -3_950_000 < y < -2_400_000 for x, y in pts)  # Lo29 CAD coordinates


def test_report_carries_the_results_and_stamp_on_every_page(rendered):
    text = pdf_text(data(rendered, "report"))
    assert "Electrification design report" in text
    assert text.count("Revision D1") >= 3  # footer on every page
    assert "NOT FOR SUBMISSION" in text and "Notified maximum demand" in text and "Traceability" in text


def test_boq_avoids_double_counting_transformers_and_totals_in_excel(rendered):
    secs = sections(package())
    assert not any(x["item"].startswith("Transformer") for title, lines, _ in secs if title.startswith("LV") for x in lines)
    assert any("pole-mount" in x["item"] for title, lines, _ in secs if title.startswith("MV") for x in lines)
    wb = openpyxl.load_workbook(io.BytesIO(data(rendered, "boq_xlsx")))
    ws = wb["BoQ"]
    assert "ESTIMATE" in ws["A2"].value and ws["A6"].value == "Revision D1"
    assert str(ws.cell(ws.max_row, 6).value).startswith("=SUM(F")
    assert {"Summary", "Assemblies", "Stamp"} <= set(wb.sheetnames)
    assert "ESTIMATE" in pdf_text(data(rendered, "boq_pdf"))


def test_gis_exports(rendered):
    g = json.loads(data(rendered, "gis_geojson"))
    assert g["reticula"]["revision"] == "D1"
    assert {f["properties"]["layer"] for f in g["features"]} == {"lv_branches", "lv_nodes", "customers", "mv_branches", "mv_nodes"}
    k = minidom.parseString(data(rendered, "gis_kml"))
    assert "Revision D1" in k.getElementsByTagName("description")[0].firstChild.data
    z = zipfile.ZipFile(io.BytesIO(data(rendered, "gis_shp")))
    r = shapefile.Reader(shp=io.BytesIO(z.read("customers.shp")), shx=io.BytesIO(z.read("customers.shx")), dbf=io.BytesIO(z.read("customers.dbf")))
    assert len(r) == 48 and r.shapeType == shapefile.POINT
    assert "RST" in [rec["phases"] for rec in r.records()] or "RWB" in [rec["phases"] for rec in r.records()]
    assert "Revision D1" in z.read("README.txt").decode()


def test_load_schedule_spreadsheet(rendered):
    ws = openpyxl.load_workbook(io.BytesIO(data(rendered, "loads_xlsx"))).active
    rows = list(ws.iter_rows(values_only=True))
    header = rows.index(next(r for r in rows if r and r[0] == "Erf"))
    assert len(rows) - header - 2 == 48 and rows[-1][0] == "Total"


def test_checklist_reports_what_is_met(rendered):
    status = {i.id: i.status for i in rendered.checklist}
    assert status["C01"] == "met" and status["C03"] == "met"
    assert status["C06"] == "not met"  # unverified rules values
    assert status["C07"] == "not met"  # one open assumption
    assert status["C10"] == "manual"


def test_submission_pack_holds_every_document_and_the_registers(rendered):
    z = zipfile.ZipFile(io.BytesIO(data(rendered, "submission_pack")))
    names = set(z.namelist())
    assert {f.name for f in rendered.files if f.kind != "submission_pack"} <= names
    assert "tt1-d1-registers.csv" in names and "STAMP.txt" in names
    reg = z.read("tt1-d1-registers.csv").decode()
    assert "RET-TT1-001" in reg and "checklist,C06" in reg


def test_endpoint_returns_only_the_requested_kinds(client):
    res = client.post("/docs/render", json={"package": package().model_dump(mode="json"), "kinds": ["gis_geojson", "loads_xlsx"]})
    assert res.status_code == 200, res.text
    assert [f["kind"] for f in res.json()["files"]] == ["gis_geojson", "loads_xlsx"]
